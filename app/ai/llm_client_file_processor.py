from __future__ import annotations

import base64
import io
import mimetypes
from pathlib import Path
from typing import Any


IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}
TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".json",
    ".yaml",
    ".yml",
    ".csv",
    ".log",
    ".ini",
    ".toml",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".go",
    ".rs",
    ".cpp",
    ".c",
    ".h",
    ".hpp",
    ".sql",
    ".xml",
    ".html",
    ".css",
}


class FileProcessor:
    def __init__(self, config: Any) -> None:
        self.config = config

    def safe_read_text(self, path: Path) -> str:
        for encoding in ("utf-8", "gb18030", "utf-16"):
            try:
                return path.read_text(encoding=encoding, errors="strict")
            except Exception:
                continue
        return path.read_text(encoding="utf-8", errors="replace")

    def extract_pdf_text(self, path: Path) -> str:
        try:
            from pypdf import PdfReader  # type: ignore
        except Exception:
            return "[PDF parser unavailable: install `pypdf`]"
        try:
            reader = PdfReader(str(path))
            pages: list[str] = []
            for page in reader.pages[:30]:
                text = (page.extract_text() or "").strip()
                if text:
                    pages.append(text)
            return "\n\n".join(pages).strip()
        except Exception as exc:  # noqa: BLE001
            return f"[PDF read failed: {exc}]"

    def extract_docx_text(self, path: Path) -> str:
        try:
            from docx import Document  # type: ignore
        except Exception:
            return "[DOCX parser unavailable: install `python-docx`]"
        try:
            document = Document(str(path))
            lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            return "\n".join(lines).strip()
        except Exception as exc:  # noqa: BLE001
            return f"[DOCX read failed: {exc}]"

    def extract_xlsx_text(self, path: Path) -> str:
        try:
            from openpyxl import load_workbook  # type: ignore
        except Exception:
            return "[XLSX parser unavailable: install `openpyxl`]"
        try:
            wb = load_workbook(filename=str(path), read_only=True, data_only=True)
            rows_out: list[str] = []
            for ws in wb.worksheets[:5]:
                rows_out.append(f"# Sheet: {ws.title}")
                for idx, row in enumerate(ws.iter_rows(values_only=True)):
                    if idx >= 200:
                        rows_out.append("[...truncated rows...]")
                        break
                    values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if values:
                        rows_out.append(" | ".join(values))
            wb.close()
            return "\n".join(rows_out).strip()
        except Exception as exc:  # noqa: BLE001
            return f"[XLSX read failed: {exc}]"

    def extract_file_text(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in TEXT_SUFFIXES:
            return self.safe_read_text(path)
        if suffix == ".pdf":
            return self.extract_pdf_text(path)
        if suffix == ".docx":
            return self.extract_docx_text(path)
        if suffix in {".xlsx", ".xlsm"}:
            return self.extract_xlsx_text(path)
        return "[Unsupported file format for text extraction]"

    def make_image_block(self, path: Path, force_safe_transcode: bool = False) -> dict[str, Any]:
        raw_size = path.stat().st_size
        if raw_size > self.config.max_image_bytes * 4:
            size_mb = raw_size / (1024 * 1024)
            limit_mb = (self.config.max_image_bytes * 4) / (1024 * 1024)
            raise ValueError(f"image too large: {size_mb:.2f}MB > {limit_mb:.2f}MB ({path.name})")

        try:
            from PIL import Image, ImageOps  # type: ignore
        except Exception:
            size = path.stat().st_size
            if size > self.config.max_image_bytes:
                size_mb = size / (1024 * 1024)
                limit_mb = self.config.max_image_bytes / (1024 * 1024)
                raise ValueError(
                    f"image too large without Pillow: {size_mb:.2f}MB > {limit_mb:.2f}MB ({path.name})"
                )
            mime_type = mimetypes.guess_type(path.name)[0] or "image/png"
            encoded = base64.b64encode(path.read_bytes()).decode("ascii")
            data_url = f"data:{mime_type};base64,{encoded}"
            return {"type": "image_url", "image_url": {"url": data_url}}

        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image)
            original_mime = mimetypes.guess_type(path.name)[0] or "image/png"
            original_max_side = max(image.size)
            if (
                not force_safe_transcode
                and (
                    original_mime == "image/png"
                    and raw_size <= self.config.max_image_bytes
                    and original_max_side <= self.config.max_image_side
                )
            ):
                encoded = base64.b64encode(path.read_bytes()).decode("ascii")
                data_url = f"data:{original_mime};base64,{encoded}"
                return {"type": "image_url", "image_url": {"url": data_url}}

            if image.mode not in ("RGB", "RGBA"):
                image = image.convert("RGBA" if "A" in image.getbands() else "RGB")

            max_side = max(image.size)
            if max_side > self.config.max_image_side:
                scale = self.config.max_image_side / max_side
                new_size = (
                    max(1, int(image.size[0] * scale)),
                    max(1, int(image.size[1] * scale)),
                )
                image = image.resize(new_size, Image.Resampling.LANCZOS)

            buffer = io.BytesIO()
            if force_safe_transcode:
                safe_image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
                safe_max_side = max(safe_image.size)
                if safe_max_side > 768:
                    scale = 768 / safe_max_side
                    new_size = (
                        max(1, int(safe_image.size[0] * scale)),
                        max(1, int(safe_image.size[1] * scale)),
                    )
                    safe_image = safe_image.resize(new_size, Image.Resampling.LANCZOS)
                while True:
                    buffer = io.BytesIO()
                    safe_image.save(buffer, format="PNG", optimize=False)
                    payload = buffer.getvalue()
                    if len(payload) <= self.config.max_image_bytes:
                        mime_type = "image/png"
                        break
                    max_side = max(safe_image.size)
                    if max_side <= 512:
                        raise ValueError(
                            f"safe PNG image still too large: {len(payload) / (1024 * 1024):.2f}MB ({path.name})"
                        )
                    scale = 0.8
                    new_size = (
                        max(1, int(safe_image.size[0] * scale)),
                        max(1, int(safe_image.size[1] * scale)),
                    )
                    safe_image = safe_image.resize(new_size, Image.Resampling.LANCZOS)
            elif "A" in image.getbands():
                image.save(buffer, format="PNG", optimize=False)
                mime_type = "image/png"
            else:
                image = image.convert("RGB")
                image.save(
                    buffer,
                    format="JPEG",
                    quality=self.config.image_jpeg_quality,
                    optimize=False,
                    progressive=False,
                )
                mime_type = "image/jpeg"

        payload = buffer.getvalue()
        if len(payload) > self.config.max_image_bytes:
            raise ValueError(
                f"optimized image still too large: {len(payload) / (1024 * 1024):.2f}MB ({path.name})"
            )

        encoded = base64.b64encode(payload).decode("ascii")
        data_url = f"data:{mime_type};base64,{encoded}"
        return {"type": "image_url", "image_url": {"url": data_url}}

    def normalize_attachments(self, attachment_paths: Iterable[str | Path] | None) -> list[Path]:
        if not attachment_paths:
            return []
        out: list[Path] = []
        seen: set[Path] = set()
        for raw in attachment_paths:
            path = Path(raw).expanduser()
            try:
                path = path.resolve()
            except Exception:
                path = path.absolute()
            if path in seen:
                continue
            seen.add(path)
            out.append(path)
        return out[: self.config.max_attachment_files]

    def build_user_message_content(
        self,
        user_input: str,
        attachment_paths: Iterable[str | Path] | None,
        force_safe_transcode: bool = False,
    ) -> tuple[str | list[dict[str, Any]], list[str], bool]:
        text_sections: list[str] = []
        images: list[dict[str, Any]] = []
        notices: list[str] = []
        total_chars = 0

        prompt_text = user_input.strip()
        if prompt_text:
            text_sections.append(prompt_text)

        for path in self.normalize_attachments(attachment_paths):
            if not path.exists():
                notices.append(f"Attachment not found: {path}")
                continue
            if path.is_dir():
                notices.append(f"Attachment is a directory and was skipped: {path.name}")
                continue

            suffix = path.suffix.lower()
            if suffix in IMAGE_SUFFIXES and self.config.enable_multimodal:
                try:
                    images.append(self.make_image_block(path, force_safe_transcode=force_safe_transcode))
                except Exception as exc:  # noqa: BLE001
                    notices.append(f"Image skipped ({path.name}): {exc}")
                continue

            extracted = self.extract_file_text(path).strip()
            if not extracted:
                notices.append(f"File has no readable text: {path.name}")
                continue

            remaining_total = max(0, self.config.max_attachment_chars_total - total_chars)
            if remaining_total == 0:
                notices.append("Attachment text budget reached; remaining files skipped.")
                break

            max_for_file = min(self.config.max_attachment_chars_per_file, remaining_total)
            clipped = extracted[:max_for_file]
            total_chars += len(clipped)
            if len(clipped) < len(extracted):
                clipped += "\n...[truncated]..."

            text_sections.append(f"[File: {path.name}]\n{clipped}")

        if notices:
            text_sections.append("Attachment notes:\n" + "\n".join(f"- {notice}" for notice in notices))

        merged_text = "\n\n".join(section for section in text_sections if section).strip()
        if not images:
            return merged_text or "(empty)", notices, False

        content_blocks: list[dict[str, Any]] = [{"type": "text", "text": merged_text or "Please analyze the images."}]
        content_blocks.extend(images)
        return content_blocks, notices, True
