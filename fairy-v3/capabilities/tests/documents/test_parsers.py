from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document as WordDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from fairy_capabilities.documents.parsers import (
    CompositeDocumentParser,
    DocumentParseError,
)


def _pdf_bytes(text: str, *, encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_reference})}
    )
    stream = DecodedStreamObject()
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream.set_data(f"BT /F1 12 Tf 20 200 Td ({escaped}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("fixture-password")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _docx_bytes() -> bytes:
    document = WordDocument()
    document.add_heading("Fairy Architecture", level=1)
    document.add_paragraph("Project-first document retrieval.")
    document.add_heading("Safety", level=2)
    document.add_paragraph("RAG is not Hermes memory.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "media_type", "content", "expected_parser", "expected_text"),
    [
        (
            "notes.txt",
            "text/plain",
            b"Fairy notes\r\nProject-first.",
            "plain_text",
            "Project-first.",
        ),
        (
            "architecture.md",
            "text/markdown",
            b"# Fairy\n\nProject-first.\n\n## Safety\nNo host shell.",
            "markdown",
            "No host shell.",
        ),
        (
            "page.html",
            "text/html",
            b"<html><head><title>Fairy</title><script>ignore()</script></head>"
            b"<body><h1>Architecture</h1><p>Command Bus only.</p></body></html>",
            "html",
            "Command Bus only.",
        ),
        (
            "paper.pdf",
            "application/pdf",
            _pdf_bytes("Evidence from PDF"),
            "pypdf",
            "Evidence from PDF",
        ),
        (
            "spec.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _docx_bytes(),
            "python_docx",
            "RAG is not Hermes memory.",
        ),
    ],
    ids=("plain", "markdown", "html", "pdf", "docx"),
)
def test_supported_parsers_return_bounded_sections_with_locator_provenance(
    filename: str,
    media_type: str,
    content: bytes,
    expected_parser: str,
    expected_text: str,
) -> None:
    extracted = CompositeDocumentParser().parse(
        filename=filename,
        media_type=media_type,
        content=content,
    )

    assert extracted.media_type == media_type
    assert extracted.parser == expected_parser
    assert extracted.parser_version
    assert [section.ordinal for section in extracted.sections] == list(
        range(len(extracted.sections))
    )
    assert expected_text in "\n".join(section.text for section in extracted.sections)
    assert all(section.locator for section in extracted.sections)
    assert "ignore()" not in "\n".join(section.text for section in extracted.sections)


@pytest.mark.parametrize(
    ("filename", "media_type", "content", "expected_code"),
    [
        ("secret.pdf", "application/pdf", _pdf_bytes("Secret", encrypted=True), "ENCRYPTED"),
        ("broken.pdf", "application/pdf", b"%PDF-not-valid", "MALFORMED"),
        ("image.png", "image/png", b"png", "UNSUPPORTED_MEDIA_TYPE"),
        ("wrong.pdf", "text/plain", b"not a pdf", "MEDIA_TYPE_MISMATCH"),
        ("bad.txt", "text/plain", b"\xff\xfe\x00", "MALFORMED"),
    ],
    ids=("encrypted-pdf", "malformed-pdf", "unsupported", "mismatch", "invalid-utf8"),
)
def test_parser_rejects_encrypted_malformed_unsupported_and_mismatched_inputs(
    filename: str,
    media_type: str,
    content: bytes,
    expected_code: str,
) -> None:
    with pytest.raises(DocumentParseError) as captured:
        CompositeDocumentParser().parse(
            filename=filename,
            media_type=media_type,
            content=content,
        )

    assert captured.value.error_code == f"DOCUMENT_{expected_code}"


def test_parser_rejects_input_and_extracted_text_above_limits() -> None:
    with pytest.raises(DocumentParseError) as input_error:
        CompositeDocumentParser(max_input_bytes=8).parse(
            filename="large.txt",
            media_type="text/plain",
            content=b"123456789",
        )
    with pytest.raises(DocumentParseError) as text_error:
        CompositeDocumentParser(max_extracted_characters=10).parse(
            filename="large.txt",
            media_type="text/plain",
            content=b"more than ten characters",
        )

    assert input_error.value.error_code == "DOCUMENT_TOO_LARGE"
    assert text_error.value.error_code == "DOCUMENT_TEXT_TOO_LARGE"
