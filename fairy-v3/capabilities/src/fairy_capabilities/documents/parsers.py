from __future__ import annotations

import re
import zipfile
from html.parser import HTMLParser
from importlib.metadata import version
from io import BytesIO
from pathlib import PurePath

from docx import Document as WordDocument
from fairy_core.documents import ExtractedDocument, ExtractedSection
from fairy_core.documents.models import normalized_filename
from pypdf import PdfReader
from pypdf.errors import PdfReadError

_DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_EXTENSIONS = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".pdf": "application/pdf",
    ".docx": _DOCX_MEDIA_TYPE,
}
_WHITESPACE = re.compile(r"[ \t\f\v]+")


class DocumentParseError(ValueError):
    def __init__(self, error_code: str, message: str) -> None:
        self.error_code = error_code
        super().__init__(message)


class CompositeDocumentParser:
    def __init__(
        self,
        *,
        max_input_bytes: int = 20 * 1024 * 1024,
        max_extracted_characters: int = 5_000_000,
        max_sections: int = 10_000,
        max_pdf_pages: int = 500,
        max_container_uncompressed_bytes: int = 100 * 1024 * 1024,
    ) -> None:
        limits = (
            max_input_bytes,
            max_extracted_characters,
            max_sections,
            max_pdf_pages,
            max_container_uncompressed_bytes,
        )
        if any(isinstance(value, bool) or value < 1 for value in limits):
            raise ValueError("document parser limits must be positive integers")
        self._max_input_bytes = max_input_bytes
        self._max_extracted_characters = max_extracted_characters
        self._max_sections = max_sections
        self._max_pdf_pages = max_pdf_pages
        self._max_container_uncompressed_bytes = max_container_uncompressed_bytes

    def parse(
        self,
        *,
        filename: str,
        media_type: str,
        content: bytes,
    ) -> ExtractedDocument:
        try:
            safe_filename = normalized_filename(filename)
        except ValueError as error:
            raise DocumentParseError("DOCUMENT_INVALID_FILENAME", str(error)) from error
        if not isinstance(content, bytes):
            raise DocumentParseError("DOCUMENT_MALFORMED", "document content must be bytes")
        if not content:
            raise DocumentParseError("DOCUMENT_MALFORMED", "document content is empty")
        if len(content) > self._max_input_bytes:
            raise DocumentParseError("DOCUMENT_TOO_LARGE", "document exceeds the input limit")
        normalized_media = media_type.strip().lower()
        expected_media = _EXTENSIONS.get(PurePath(safe_filename).suffix.casefold())
        if expected_media is None or normalized_media not in set(_EXTENSIONS.values()):
            raise DocumentParseError(
                "DOCUMENT_UNSUPPORTED_MEDIA_TYPE",
                "document media type is unsupported",
            )
        if normalized_media != expected_media:
            raise DocumentParseError(
                "DOCUMENT_MEDIA_TYPE_MISMATCH",
                "document media type does not match its filename",
            )
        try:
            if normalized_media == "text/plain":
                parser, parser_version, sections = "plain_text", "1", self._plain(content)
            elif normalized_media == "text/markdown":
                parser, parser_version, sections = "markdown", "1", self._markdown(content)
            elif normalized_media == "text/html":
                parser, parser_version, sections = "html", "1", self._html(content)
            elif normalized_media == "application/pdf":
                parser, parser_version, sections = "pypdf", version("pypdf"), self._pdf(content)
            else:
                parser, parser_version, sections = (
                    "python_docx",
                    version("python-docx"),
                    self._docx(content),
                )
        except DocumentParseError:
            raise
        except Exception as error:
            raise DocumentParseError(
                "DOCUMENT_MALFORMED",
                "document parser rejected malformed content",
            ) from error
        finalized = tuple(section for section in sections if section.text.strip())
        if not finalized:
            raise DocumentParseError(
                "DOCUMENT_NO_TEXT",
                "document contains no extractable text",
            )
        if len(finalized) > self._max_sections:
            raise DocumentParseError(
                "DOCUMENT_TOO_MANY_SECTIONS",
                "document exceeds the section limit",
            )
        total_characters = sum(len(section.text) for section in finalized)
        if total_characters > self._max_extracted_characters:
            raise DocumentParseError(
                "DOCUMENT_TEXT_TOO_LARGE",
                "document extracted text exceeds the limit",
            )
        normalized_sections = tuple(
            ExtractedSection(
                ordinal=ordinal,
                title=section.title,
                text=section.text,
                locator=section.locator,
            )
            for ordinal, section in enumerate(finalized)
        )
        return ExtractedDocument(
            media_type=normalized_media,
            parser=parser,
            parser_version=parser_version,
            sections=normalized_sections,
        )

    def _plain(self, content: bytes) -> tuple[ExtractedSection, ...]:
        text = _decode_utf8(content)
        return (
            ExtractedSection(
                ordinal=0,
                title="Document",
                text=text,
                locator={"kind": "text_document"},
            ),
        )

    def _markdown(self, content: bytes) -> tuple[ExtractedSection, ...]:
        text = _decode_utf8(content)
        sections: list[ExtractedSection] = []
        title = "Document"
        lines: list[str] = []
        heading_line = 1
        for line_number, line in enumerate(text.splitlines(), start=1):
            match = re.match(r"^#{1,6}\s+(.+?)\s*#*\s*$", line)
            if match:
                if "\n".join(lines).strip():
                    sections.append(
                        ExtractedSection(
                            ordinal=len(sections),
                            title=title,
                            text="\n".join(lines),
                            locator={"heading": title, "line": heading_line},
                        )
                    )
                title = match.group(1).strip()
                lines = []
                heading_line = line_number
            else:
                lines.append(line)
        if "\n".join(lines).strip():
            sections.append(
                ExtractedSection(
                    ordinal=len(sections),
                    title=title,
                    text="\n".join(lines),
                    locator={"heading": title, "line": heading_line},
                )
            )
        return tuple(sections)

    def _html(self, content: bytes) -> tuple[ExtractedSection, ...]:
        parser = _SafeHtmlTextParser()
        parser.feed(_decode_utf8(content))
        parser.close()
        title = _normalize_inline(" ".join(parser.title)) or "HTML document"
        text = _normalize_blocks(parser.blocks)
        return (
            ExtractedSection(
                ordinal=0,
                title=title,
                text=text,
                locator={"kind": "html_document"},
            ),
        )

    def _pdf(self, content: bytes) -> tuple[ExtractedSection, ...]:
        try:
            reader = PdfReader(BytesIO(content), strict=True)
        except PdfReadError as error:
            raise DocumentParseError("DOCUMENT_MALFORMED", "PDF is malformed") from error
        if reader.is_encrypted:
            raise DocumentParseError("DOCUMENT_ENCRYPTED", "encrypted PDF is not supported")
        if len(reader.pages) > self._max_pdf_pages:
            raise DocumentParseError("DOCUMENT_TOO_MANY_SECTIONS", "PDF exceeds the page limit")
        sections: list[ExtractedSection] = []
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as error:
                raise DocumentParseError(
                    "DOCUMENT_MALFORMED", "PDF text extraction failed"
                ) from error
            if text.strip():
                sections.append(
                    ExtractedSection(
                        ordinal=len(sections),
                        title=f"Page {page_number}",
                        text=_normalize_text(text),
                        locator={"page": page_number},
                    )
                )
        return tuple(sections)

    def _docx(self, content: bytes) -> tuple[ExtractedSection, ...]:
        self._validate_zip(content)
        try:
            document = WordDocument(BytesIO(content))
        except Exception as error:
            raise DocumentParseError("DOCUMENT_MALFORMED", "DOCX is malformed") from error
        sections: list[ExtractedSection] = []
        title = "Document"
        paragraphs: list[str] = []
        start_paragraph = 1
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            style_name = (paragraph.style.name or "").casefold() if paragraph.style else ""
            if text and style_name.startswith("heading"):
                if paragraphs:
                    sections.append(
                        ExtractedSection(
                            ordinal=len(sections),
                            title=title,
                            text="\n".join(paragraphs),
                            locator={"heading": title, "paragraph": start_paragraph},
                        )
                    )
                title = text
                paragraphs = []
                start_paragraph = index
            elif text:
                paragraphs.append(text)
        if paragraphs:
            sections.append(
                ExtractedSection(
                    ordinal=len(sections),
                    title=title,
                    text="\n".join(paragraphs),
                    locator={"heading": title, "paragraph": start_paragraph},
                )
            )
        return tuple(sections)

    def _validate_zip(self, content: bytes) -> None:
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                members = archive.infolist()
        except (OSError, zipfile.BadZipFile) as error:
            raise DocumentParseError("DOCUMENT_MALFORMED", "DOCX container is malformed") from error
        if len(members) > 10_000:
            raise DocumentParseError("DOCUMENT_TOO_LARGE", "DOCX contains too many members")
        total_size = 0
        for member in members:
            path = PurePath(member.filename)
            if path.is_absolute() or ".." in path.parts:
                raise DocumentParseError("DOCUMENT_MALFORMED", "DOCX contains an unsafe path")
            total_size += member.file_size
            if member.file_size > 0 and member.compress_size == 0:
                raise DocumentParseError("DOCUMENT_TOO_LARGE", "DOCX compression is invalid")
            if member.compress_size > 0 and member.file_size / member.compress_size > 200:
                raise DocumentParseError("DOCUMENT_TOO_LARGE", "DOCX compression ratio is unsafe")
        if total_size > self._max_container_uncompressed_bytes:
            raise DocumentParseError("DOCUMENT_TOO_LARGE", "DOCX expands beyond the limit")


class _SafeHtmlTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[str] = []
        self.title: list[str] = []
        self._ignored_depth = 0
        self._title_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        del attrs
        normalized = tag.casefold()
        if normalized in {"script", "style", "noscript", "template"}:
            self._ignored_depth += 1
        if normalized == "title":
            self._title_depth += 1

    def handle_endtag(self, tag: str) -> None:
        normalized = tag.casefold()
        if normalized in {"script", "style", "noscript", "template"}:
            self._ignored_depth = max(0, self._ignored_depth - 1)
        if normalized == "title":
            self._title_depth = max(0, self._title_depth - 1)

    def handle_data(self, data: str) -> None:
        if self._ignored_depth:
            return
        normalized = _normalize_inline(data)
        if normalized:
            self.blocks.append(normalized)
            if self._title_depth:
                self.title.append(normalized)


def _decode_utf8(content: bytes) -> str:
    try:
        text = content.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as error:
        raise DocumentParseError("DOCUMENT_MALFORMED", "text document must be UTF-8") from error
    if "\x00" in text:
        raise DocumentParseError("DOCUMENT_MALFORMED", "text document contains NUL bytes")
    return _normalize_text(text)


def _normalize_text(value: str) -> str:
    lines = [
        _WHITESPACE.sub(" ", line).strip()
        for line in value.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    ]
    return "\n".join(lines).strip()


def _normalize_inline(value: str) -> str:
    return _WHITESPACE.sub(" ", value).strip()


def _normalize_blocks(values: list[str]) -> str:
    return "\n".join(value for value in values if value).strip()


__all__ = ["CompositeDocumentParser", "DocumentParseError"]
